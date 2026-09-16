from pathlib import Path
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

csv_path = Path("tensorboard_scalars.csv")
out = Path("/mnt/c/Users/orori/PycharmProjects/ESPnet_Project/Report/ESPnet_Experiment_Report_Final.docx")
chart1 = Path("/mnt/c/Users/orori/PycharmProjects/ESPnet_Project/Report/validation_loss_by_epoch_final.png")
chart2 = Path("/mnt/c/Users/orori/PycharmProjects/ESPnet_Project/Report/validation_ctc_cer_by_epoch_final.png")

df = pd.read_csv(csv_path)
df["split"] = df["event_file"].str.extract(r"/tensorboard/(train|valid)/")[0]
df["epoch"] = (df["step"] / 2191).round().astype(int)

def valid_table(exp, event_file=None):
    sub = df[(df["experiment"] == exp) & (df["split"] == "valid")].copy()
    if event_file is not None:
        sub = sub[sub["event_file"] == event_file]
    return sub.pivot_table(index="epoch", columns="tag", values="value", aggfunc="first").sort_index()

# For lr5e5_ctc03, use only the later/final training session.
ctc_exp = "hebrew_xlsr_finetune_lr5e5_ctc03"
ctc_valid_files = list(df[(df["experiment"] == ctc_exp) & (df["split"] == "valid")]["event_file"].drop_duplicates())
final_ctc03_file = [f for f in ctc_valid_files if "1788728361" in f][0]

series = {
    "finetune (10 ep)": valid_table("hebrew_xlsr_finetune"),
    "finetune_30ep": valid_table("hebrew_xlsr_finetune_30ep"),
    "finetune_lr5e5": valid_table("hebrew_xlsr_finetune_lr5e5"),
    "finetune_lr5e5_ctc03": valid_table(ctc_exp, final_ctc03_file),
}

# Validation loss chart
plt.figure(figsize=(7.2, 4.2))
for label, t in series.items():
    plt.plot(t.index, t["loss"], marker="o", markersize=2.5, label=label)
plt.xlabel("Epoch")
plt.ylabel("Validation loss")
plt.title("Validation Loss by Epoch")
plt.legend(fontsize=7)
plt.tight_layout()
plt.savefig(chart1, dpi=170)
plt.close()

# Validation CTC CER chart
plt.figure(figsize=(7.2, 4.2))
for label, t in series.items():
    plt.plot(t.index, t["cer_ctc"], marker="o", markersize=2.5, label=label)
plt.xlabel("Epoch")
plt.ylabel("Validation CTC CER")
plt.title("Validation CTC CER by Epoch")
plt.legend(fontsize=7)
plt.tight_layout()
plt.savefig(chart2, dpi=170)
plt.close()

# Summaries
summary_rows = []
for label, t in series.items():
    best_loss_ep = int(t["loss"].idxmin())
    best_acc_ep = int(t["acc"].idxmax())
    best_ctc_ep = int(t["cer_ctc"].idxmin())
    best_att_cer_ep = int(t["cer"].idxmin())
    best_wer_ep = int(t["wer"].idxmin())
    summary_rows.append({
        "run": label,
        "epochs": len(t),
        "best_loss": float(t.loc[best_loss_ep, "loss"]),
        "best_loss_ep": best_loss_ep,
        "best_acc": float(t.loc[best_acc_ep, "acc"]),
        "best_acc_ep": best_acc_ep,
        "best_ctc_cer": float(t.loc[best_ctc_ep, "cer_ctc"]),
        "best_ctc_ep": best_ctc_ep,
        "best_att_cer": float(t.loc[best_att_cer_ep, "cer"]),
        "best_att_cer_ep": best_att_cer_ep,
        "best_wer": float(t.loc[best_wer_ep, "wer"]),
        "best_wer_ep": best_wer_ep,
    })

my_test = [
    (1, "decode_my_test_ctc", "2026-09-05 17:56:47", 0.246, 0.854, 0.246),
    (2, "decode_my_test_ctc03", "2026-09-05 18:49:45", 0.409, 1.415, 0.379),
    (3, "decode_my_test_ctc05", "2026-09-05 18:55:09", 0.313, 1.037, 0.315),
    (4, "decode_my_test_ctc07", "2026-09-05 20:53:36", 0.299, 0.967, 0.303),
    (5, "decode_my_test_ctc09", "2026-09-05 21:19:13", 0.301, 0.951, 0.303),
    (6, "decode_my_test_ctc08", "2026-09-05 21:51:54", 0.301, 0.967, 0.303),
    (7, "decode_my_test_att", "2026-09-05 22:01:31", 4.894, 13.862, 3.034),
    (8, "decode_my_test_hybrid03_lr5e5_ctc03", "2026-09-07 13:19:07", 0.162, 0.573, 0.162),
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)

for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
    doc.styles[style_name].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Title"].font.size = Pt(22)
doc.styles["Heading 1"].font.size = Pt(16)
doc.styles["Heading 2"].font.size = Pt(12.5)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("ESPnet Hebrew ASR Experiments\nFine-Tuning, Decoding, Results and Discussion")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Final version using the confirmed fine-tuning configuration and TensorBoard results")
r.bold = True

doc.add_heading("1. Experimental structure", level=1)
doc.add_paragraph(
    "The project separates model training from evaluation. The pretrained XLS-R 300M frontend is fine-tuned on the Hebrew Campus training split. "
    "The Hebrew Campus dev split is used as validation data during training. The custom my_test set is used afterward for decoding experiments and "
    "does not update model parameters."
)

doc.add_heading("2. Common model and data setup", level=1)
for x in [
    "Training data: hebrew_campus_subset/data/train.",
    "Validation data: hebrew_campus_subset/data/dev.",
    "Frontend: S3PRL XLS-R 300M at 16 kHz.",
    "Preencoder: linear, 1024 -> 256.",
    "Encoder: 4-block Transformer, output size 256, 4 attention heads, 1024 linear units.",
    "Decoder: 2-block Transformer, 4 attention heads, 1024 linear units.",
    "Tokenization: Hebrew character tokens.",
    "Optimizer: AdamW with weight decay 0.01.",
    "Batch size: 1 with gradient accumulation of 8; mixed precision enabled.",
]:
    doc.add_paragraph(x, style="List Bullet")

doc.add_heading("3. Fine-tuning experiments: what changed", level=1)
tbl = doc.add_table(rows=1, cols=6)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["Experiment", "Max epochs", "Patience", "Learning rate", "Training CTC weight", "Observed epochs"]):
    tbl.rows[0].cells[i].text = h
for row in [
    ("hebrew_xlsr_finetune", "10", "3", "1e-5", "0.5", "10"),
    ("hebrew_xlsr_finetune_30ep", "30", "7", "1e-5", "0.5", "20"),
    ("hebrew_xlsr_finetune_lr5e5", "30", "7", "5e-5", "0.5", "16"),
    ("hebrew_xlsr_finetune_lr5e5_ctc03", "30", "7", "5e-5", "0.3", "30"),
]:
    cells = tbl.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v

doc.add_paragraph(
    "The progression was controlled: first increase the training horizon, then increase the learning rate, then reduce the training CTC weight "
    "from 0.5 to 0.3 so the attention objective receives relatively more influence during training."
)

doc.add_heading("4. Note about the lr5e5_ctc03 TensorBoard logs", level=1)
doc.add_paragraph(
    "The TensorBoard directory contains logs from two training sessions with the same lr5e5_ctc03 configuration. These are not two different experiments. "
    "For the figures and comparisons in this report, only the later/final session is used. The checkpoint files in the experiment directory show that the "
    "final training state selected epoch 29 as both valid.loss.best.pth and valid.acc.best.pth, while latest.pth points to epoch 30."
)

doc.add_heading("5. Fine-tuning results from validation data", level=1)
tbl = doc.add_table(rows=1, cols=7)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["Run", "Epochs", "Best valid loss", "Best valid acc", "Best CTC CER", "Best att. CER", "Best att. WER"]):
    tbl.rows[0].cells[i].text = h
for s in summary_rows:
    vals = [
        s["run"],
        str(s["epochs"]),
        f'{s["best_loss"]:.3f} (ep {s["best_loss_ep"]})',
        f'{s["best_acc"]:.3f} (ep {s["best_acc_ep"]})',
        f'{s["best_ctc_cer"]:.3f} (ep {s["best_ctc_ep"]})',
        f'{s["best_att_cer"]:.3f} (ep {s["best_att_cer_ep"]})',
        f'{s["best_wer"]:.3f} (ep {s["best_wer_ep"]})',
    ]
    cells = tbl.add_row().cells
    for i, v in enumerate(vals):
        cells[i].text = v

doc.add_picture(str(chart1), width=Inches(6.7))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_picture(str(chart2), width=Inches(6.7))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("6. Fine-tuning experiment discussion", level=1)

doc.add_heading("6.1 hebrew_xlsr_finetune: initial 10-epoch run", level=2)
doc.add_paragraph(
    "This run used lr=1e-5 and training ctc_weight=0.5. Validation loss continued falling through epoch 10, reaching 105.032, while CTC CER improved strongly "
    "to 0.230. The attention-side CER remained around 0.76 and WER remained above 1.4. The CTC branch therefore learned useful acoustic/alignment behavior much "
    "faster than the attention branch. Because the run ended at its configured maximum while validation loss was still decreasing, extending the training horizon "
    "was a reasonable next experiment."
)

doc.add_heading("6.2 hebrew_xlsr_finetune_30ep: longer training", level=2)
doc.add_paragraph(
    "The main change was extending max_epoch from 10 to 30 and patience from 3 to 7. Training stopped after 20 epochs. Validation loss reached its minimum at "
    "epoch 12 (102.609), then increased, while CTC CER continued improving to 0.166 by epoch 20. The attention CER stayed near 0.75-0.77. This showed that simply "
    "training longer mainly strengthened CTC and did not solve the weak attention decoder."
)

doc.add_heading("6.3 hebrew_xlsr_finetune_lr5e5: higher learning rate", level=2)
doc.add_paragraph(
    "The learning rate was increased from 1e-5 to 5e-5 while ctc_weight remained 0.5. The best validation loss, 94.269, was reached at epoch 8, and training stopped "
    "at epoch 16. CTC CER improved to 0.107, validation attention accuracy reached 0.324, and attention CER improved to 0.733. The higher learning rate therefore "
    "accelerated optimization and modestly improved the attention branch."
)

doc.add_heading("6.4 hebrew_xlsr_finetune_lr5e5_ctc03: lower training CTC weight", level=2)
doc.add_paragraph(
    "This experiment kept lr=5e-5 but changed the training ctc_weight from 0.5 to 0.3, giving relatively more weight to the attention loss. In the final training "
    "session, validation loss reached 65.677 at epoch 29, validation accuracy reached 0.674, attention CER reached 0.390, attention WER reached 0.849, and CTC CER "
    "reached 0.085. Epoch 29 was selected as both the best validation-loss and best validation-accuracy checkpoint. This was the first configuration in the sequence "
    "where both the CTC and attention branches became substantially useful."
)

doc.add_heading("7. What the training sequence demonstrates", level=1)
for x in [
    "Extending training alone improved CTC but left the attention decoder almost unchanged.",
    "Increasing the learning rate from 1e-5 to 5e-5 improved convergence speed and modestly improved the attention branch.",
    "Reducing the training CTC weight from 0.5 to 0.3 coincided with a large improvement in attention accuracy, CER and WER while preserving strong CTC CER.",
    "The final configuration therefore produced a more balanced CTC-attention model, making it a better candidate for hybrid decoding.",
]:
    doc.add_paragraph(x, style="List Bullet")

doc.add_heading("8. my_test decoding experiments", level=1)
tbl = doc.add_table(rows=1, cols=6)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["#", "Decode experiment", "Run time", "CER", "WER", "CER no spaces"]):
    tbl.rows[0].cells[i].text = h
for row in my_test:
    cells = tbl.add_row().cells
    vals = [str(row[0]), row[1], row[2], f"{row[3]:.3f}", f"{row[4]:.3f}", f"{row[5]:.3f}"]
    for i, v in enumerate(vals):
        cells[i].text = v

doc.add_heading("9. Discussion of my_test decoding", level=1)
doc.add_paragraph(
    "The initial CTC result (CER 0.246, WER 0.854) was the strongest of the early September 5 runs. Lower CTC contributions performed worse, while the 0.7-0.9 "
    "region was more stable but still did not beat the pure-CTC baseline. Attention-only decoding failed severely (CER 4.894, WER 13.862), which is consistent with "
    "the validation evidence that the early attention branch was poorly trained."
)
doc.add_paragraph(
    "The final hybrid experiment using the later lr5e5_ctc03 training configuration achieved CER 0.162 and WER 0.573. Relative to the first CTC baseline, this is "
    "about a 34.1% relative CER reduction and a 32.9% relative WER reduction. Both the training configuration and the decoding configuration changed, so the gain "
    "should not be attributed to hybrid decoding alone."
)

doc.add_heading("10. Training CTC weight versus decoding CTC weight", level=1)
doc.add_paragraph(
    "These are different parameters. The training ctc_weight in model_conf controls how CTC loss and attention loss are combined while model parameters are learned. "
    "The decoding --ctc_weight controls how CTC and attention scores are combined during beam search after training. The lr5e5_ctc03 model used training ctc_weight=0.3; "
    "the decoding weight used in a later decode experiment is a separate setting."
)

doc.add_heading("11. Oral-exam explanation", level=1)
doc.add_paragraph(
    "A concise explanation is: 'I began with a joint CTC-attention fine-tuning setup using a 0.5 CTC training weight. The CTC branch improved quickly, but the "
    "attention branch remained weak. Extending training mostly improved CTC. I then raised the learning rate to 5e-5, which converged faster and modestly improved "
    "both branches. Finally I lowered the training CTC weight to 0.3, giving relatively more influence to the attention objective. The final run selected epoch 29 as "
    "the best validation checkpoint and substantially improved attention accuracy, CER and WER while keeping CTC CER low. On my independent my_test set, the final "
    "hybrid system reached CER 0.162 and WER 0.573.'"
)

doc.add_heading("12. Remaining reproducibility detail", level=1)
doc.add_paragraph(
    "The checkpoint directory confirms that epoch 29 was the best validation-loss and validation-accuracy checkpoint for the final lr5e5_ctc03 run. "
    "The decode folder itself does not contain the original decoding command, so the exact filename passed to --asr_model_file cannot be proven from the saved decode "
    "outputs alone. If the original shell history or decoding command is available, it can confirm whether valid.loss.best.pth, 29epoch.pth, or an averaged checkpoint "
    "was loaded."
)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(8.5)

doc.save(out)
print(f"Created: {out}")
print(f"Created: {chart1}")
print(f"Created: {chart2}")
